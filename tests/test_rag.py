import tempfile
import unittest
import uuid
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

# The application creates its Qdrant client at import time. Keep tests offline.
with patch("qdrant_client.AsyncQdrantClient"):
    from app.ai.agent.graph import ChatAgentGraph
    from app.ai.llm.provider import LLMResponse
    from app.ai.rag.agentic import EvidenceGrader
    from app.ai.rag.chunking import ChunkingService
    from app.ai.rag.retrieval import RetrievalResult, RetrievalService, RetrievedChunk
    from app.ai.rag.sparse import BM25SparseEncoder
    from app.ai.rag.text_extraction import (
        ExtractedSection,
        ExtractionResult,
        TextExtractionService,
        clean_text,
    )


USER = uuid.UUID(int=7)
KB = uuid.UUID(int=1)


class SparseEncoderTests(unittest.TestCase):
    def setUp(self):
        self.encoder = BM25SparseEncoder()

    def test_drops_stop_words_and_single_ascii_chars(self):
        self.assertEqual(self.encoder.tokenize("What is the leave policy, a?"), ["leave", "policy"])

    def test_cjk_is_indexed_as_bigrams(self):
        self.assertEqual(self.encoder.tokenize("休假政策"), ["休假", "假政", "政策"])

    def test_query_and_document_share_term_ids(self):
        doc = self.encoder.encode_document("Leave policy: employees get 20 days of leave.")
        query = self.encoder.encode_query("leave days")
        self.assertTrue(set(query.indices) <= set(doc.indices))
        self.assertEqual(query.values, [1.0, 1.0])
        # Repeated term weighs more than a single occurrence.
        weights = dict(zip(doc.indices, doc.values))
        self.assertGreater(weights[self.encoder.term_id("leave")], weights[self.encoder.term_id("days")])


class ChunkingTests(unittest.TestCase):
    def section(self, text, page=1, path="Policy"):
        return ExtractedSection(text=text, page=page, metadata={"source_type": "pdf", "section": path})

    def test_packs_paragraphs_and_spans_pages_within_section(self):
        chunker = ChunkingService(chunk_size=200, chunk_overlap=0)
        chunks = chunker.chunk_extraction(ExtractionResult([
            self.section("Alpha paragraph one.", page=1),
            self.section("Beta paragraph two.", page=2),
        ]))
        self.assertEqual(len(chunks), 1)
        self.assertEqual(chunks[0].text, "Alpha paragraph one.\n\nBeta paragraph two.")
        self.assertEqual((chunks[0].metadata["page_start"], chunks[0].metadata["page_end"]), (1, 2))

    def test_breaks_on_heading_change(self):
        chunker = ChunkingService(chunk_size=300, chunk_overlap=50, min_chunk_size=10)
        chunks = chunker.chunk_extraction(ExtractionResult([
            self.section("Eligibility rules apply to all staff.", path="Policy > Eligibility"),
            self.section("Claims must be filed in 30 days.", path="Policy > Claims"),
        ]))
        self.assertEqual([c.metadata["section"] for c in chunks], ["Policy > Eligibility", "Policy > Claims"])
        # No overlap carried across sections.
        self.assertEqual(chunks[1].metadata["overlap_chars"], 0)

    def test_overlap_is_whole_sentences_and_recorded(self):
        sentences = " ".join(f"Sentence number {i} is here." for i in range(30))
        chunker = ChunkingService(chunk_size=200, chunk_overlap=60)
        chunks = chunker.chunk_extraction(ExtractionResult([self.section(sentences)]))
        self.assertGreater(len(chunks), 2)
        for previous, chunk in zip(chunks, chunks[1:]):
            overlap = chunk.metadata["overlap_chars"]
            self.assertGreater(overlap, 0)
            carried = chunk.text[:overlap].strip()
            self.assertTrue(carried.endswith("."), carried)
            self.assertIn(carried, previous.text)
            self.assertLessEqual(len(chunk.text), 200 + 60 + 1)

    def test_streaming_feed_matches_batch(self):
        sections = [self.section(f"Paragraph {i} " + "word " * 40, page=i) for i in range(12)]
        whole = ChunkingService(chunk_size=300, chunk_overlap=40).chunk_extraction(ExtractionResult(sections))
        streaming = ChunkingService(chunk_size=300, chunk_overlap=40)
        chunks = streaming.feed(sections[:5]) + streaming.feed(sections[5:]) + streaming.flush()
        self.assertEqual([c.text for c in chunks], [c.text for c in whole])
        self.assertEqual([c.index for c in chunks], list(range(len(chunks))))

    def test_embedding_text_has_context_header(self):
        text = ChunkingService.embedding_text("3 years.", file_name="hr.pdf", section="Awards > Eligibility")
        self.assertEqual(text, "Document: hr.pdf\nSection: Awards > Eligibility\n\n3 years.")


class ExtractionTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_clean_text_fixes_ligatures_and_glyph_noise(self):
        self.assertEqual(clean_text("enƟtlement ﬁle  x"), "entitlement file x")

    async def test_pdf_headings_and_running_headers(self):
        import pymupdf as fitz

        path = self.dir / "doc.pdf"
        pdf = fitz.open()
        for number in range(1, 5):
            page = pdf.new_page()
            page.insert_text((72, 30), "ACME Corp Confidential", fontsize=9)
            page.insert_text((72, 100), f"Chapter {number}", fontsize=20)
            page.insert_text((72, 140), f"Body text for chapter {number} explains the rules.", fontsize=11)
            page.insert_text((300, 820), str(number), fontsize=9)
        pdf.save(path)
        pdf.close()

        result = await TextExtractionService().extract(path)
        text = "\n".join(s.text for s in result.sections)
        self.assertNotIn("ACME Corp Confidential", text)
        self.assertEqual(result.metadata["page_count"], 4)
        self.assertEqual(result.sections[-1].metadata["section"], "Chapter 4")
        self.assertIn("Body text for chapter 4", result.sections[-1].text)

    async def test_pdf_streams_in_page_batches(self):
        import pymupdf as fitz

        path = self.dir / "long.pdf"
        pdf = fitz.open()
        for number in range(25):
            pdf.new_page().insert_text((72, 100), f"Page {number} content about topic {number}.")
        pdf.save(path)
        pdf.close()

        batches = [b async for b in TextExtractionService().iter_batches(path, pages_per_batch=10)]
        self.assertEqual(len(batches), 3)

    async def test_docx_keeps_heading_path_and_tables(self):
        from docx import Document

        path = self.dir / "doc.docx"
        doc = Document()
        doc.add_heading("Travel Policy", level=1)
        doc.add_heading("Hotels", level=2)
        doc.add_paragraph("Stay in budget hotels.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "City", "Limit"
        table.cell(1, 0).text, table.cell(1, 1).text = "Mumbai", "5000"
        doc.save(path)

        result = await TextExtractionService().extract(path)
        last = result.sections[-1]
        self.assertEqual(last.metadata["section"], "Travel Policy > Hotels")
        self.assertIn("City | Limit\nMumbai | 5000", last.text)

    async def test_markdown_sections_and_csv_rows(self):
        md = self.dir / "a.md"
        md.write_text("# Guide\nIntro.\n## Setup\nRun it.\n```\n# not a heading\n```\n")
        sections = (await TextExtractionService().extract(md)).sections
        self.assertEqual([s.metadata["section"] for s in sections], ["Guide", "Guide > Setup"])

        csv_path = self.dir / "a.csv"
        csv_path.write_text("name,days\nCasual,12\nSick,10\n")
        text = (await TextExtractionService().extract(csv_path)).sections[0].text
        self.assertEqual(text, "name: Casual; days: 12\n\nname: Sick; days: 10")


class GraderParseTests(unittest.TestCase):
    def test_parses_fenced_and_think_wrapped_json(self):
        grade = EvidenceGrader.parse(
            '<think>hmm</think>```json\n{"relevant": ["2", 1, 9], "answerable": "Partial", '
            '"missing": "none", "next_query": "directrix equation"}\n```',
            passage_count=3,
        )
        self.assertEqual(grade.relevant, [0, 1])
        self.assertEqual(grade.answerable, "partial")
        self.assertEqual(grade.missing, "")
        self.assertEqual(grade.next_query, "directrix equation")

    def test_unparseable_returns_none(self):
        self.assertIsNone(EvidenceGrader.parse("I think passage 1 helps.", passage_count=2))


def match(index, text, *, doc=None, score=0.5, page=1, overlap=0, ingest="i1"):
    doc = doc or uuid.UUID(int=100)
    return {
        "id": f"{doc}-{index}", "score": score, "text": text, "document_id": str(doc),
        "knowledge_base_id": str(KB), "ingest_id": ingest, "chunk_index": index, "page": page,
        "page_end": page, "section": None, "file_name": "f.pdf", "metadata": {"overlap_chars": overlap},
    }


class FakeStore:
    def __init__(self, dense, lexical, neighbours=()):
        self.dense, self.lexical, self.neighbours = dense, lexical, list(neighbours)
        self.sparse_encoder = BM25SparseEncoder()
        self.searched = []

    async def hybrid_search(self, *, query, **kwargs):
        self.searched.append(query)
        return self.dense, self.lexical

    async def fetch_chunks(self, *, user_id, positions):
        return [m for m in self.neighbours if m["chunk_index"] in positions.get(m["document_id"], [])]


class RetrievalTests(unittest.IsolatedAsyncioTestCase):
    async def retrieve(self, store, **kwargs):
        with patch("app.ai.rag.retrieval.settings.RAG_NEIGHBOR_WINDOW", kwargs.pop("window", 0)):
            return await RetrievalService(vector_store=store).retrieve(
                user_id=USER, knowledge_base_ids=[KB], query=kwargs.pop("query", "error code XK-42"),
                score_threshold=0.4, **kwargs,
            )

    async def test_gate_keeps_semantic_or_exact_keyword_hits_only(self):
        doc = uuid.UUID(int=5)
        store = FakeStore(
            dense=[match(1, "Fixing timeouts in general.", score=0.62, doc=doc),
                   match(2, "Cafeteria opens at nine.", score=0.21, doc=doc)],
            lexical=[match(3, "Error code XK-42 means the token expired.", score=0.0, doc=doc)],
        )
        result = await self.retrieve(store)
        self.assertEqual({c.chunk_index for c in result.chunks}, {1, 3})
        lexical = next(c for c in result.chunks if c.chunk_index == 3)
        self.assertIsNone(lexical.metadata["dense_score"])
        self.assertEqual(lexical.metadata["lexical_coverage"], 1.0)

    async def test_runs_every_query_and_fuses(self):
        store = FakeStore(dense=[match(1, "Ellipse eccentricity.", score=0.7)], lexical=[])
        result = await self.retrieve(store, query="ellipse", queries=["directrix", "ellipse "])
        self.assertEqual(store.searched, ["ellipse", "directrix"])
        self.assertEqual(result.diagnostics["queries"], ["ellipse", "directrix"])

    async def test_per_document_cap_and_backfill(self):
        a, b = uuid.UUID(int=10), uuid.UUID(int=11)
        dense = [match(i, f"Topic alpha variant {i} " + "x" * i, doc=a, score=0.9 - i / 100) for i in range(5)]
        dense.append(match(0, "Topic alpha from another file.", doc=b, score=0.5))
        with patch("app.ai.rag.retrieval.settings.RAG_MAX_PER_DOCUMENT", 2):
            result = await self.retrieve(FakeStore(dense, []), query="topic alpha", top_k=3)
        self.assertEqual([c.document_id for c in result.chunks], [a, a, b])

    async def test_neighbour_expansion_merges_and_strips_overlap(self):
        doc = str(uuid.UUID(int=20))
        hit = match(5, "End of rule. Middle part.", doc=uuid.UUID(doc), score=0.8, page=3, overlap=13)
        neighbours = [
            match(4, "Start of rule. End of rule.", doc=uuid.UUID(doc), page=2),
            match(6, "Middle part. Final words.", doc=uuid.UUID(doc), page=3, overlap=13),
        ]
        result = await self.retrieve(FakeStore([hit], [], neighbours), query="rule", window=1)
        [passage] = result.chunks
        self.assertEqual(passage.text, "Start of rule. End of rule.\n\nMiddle part.\n\nFinal words.")
        self.assertEqual(passage.metadata["pages"], [2, 3])
        self.assertEqual(passage.metadata["chunk_indexes"], [4, 5, 6])
        self.assertEqual(passage.page, 2)


def chunk(index, text, score=0.6):
    return RetrievedChunk(
        id=f"c{index}", text=text, score=score, document_id=uuid.UUID(int=index),
        knowledge_base_id=KB, chunk_index=0, page=1, file_name="f.pdf", metadata={"pages": [1]},
    )


class AgenticLoopTests(unittest.IsolatedAsyncioTestCase):
    async def run_graph(self, rounds, grades, *, web=False, max_rounds=2):
        retrieval = SimpleNamespace(retrieve_for_bot=AsyncMock(side_effect=[
            RetrievalResult(q, chunks, [KB], diagnostics={"queries": [q]}) for q, chunks in rounds
        ]))
        llm = SimpleNamespace(chat=AsyncMock(side_effect=[
            *(LLMResponse(content=g, model="m") for g in grades),
            LLMResponse(content="final answer", model="m"),
        ]))
        tools = SimpleNamespace(execute_tool=AsyncMock(return_value={"results": []}))
        graph = ChatAgentGraph(router=None, retrieval_service=retrieval, prompt_builder=None,
                               llm_provider=llm, tool_registry=tools)
        graph._context_builder_node = AsyncMock(return_value={})
        graph._supervisor_node = AsyncMock(return_value={"route": "rag", "resolved_query": "q1"})
        generated = {}

        async def generate(state):
            generated.update(state)
            return {"response": None}

        graph._generate_node = generate
        graph.app = graph._build_graph()
        state = {"user_id": USER, "bot_id": uuid.uuid4(), "conversation_id": uuid.uuid4(),
                 "query": "q1", "has_kb": True, "db_session": object(), "model_key": "m",
                 "enable_web_search": web,
                 "available_tools": [{"function": {"name": "web_search"}}] if web else []}
        with patch("app.ai.agent.graph.settings.RAG_MAX_ROUNDS", max_rounds):
            await graph.run(state)
        return generated, retrieval, llm, tools

    async def test_refines_until_answerable_and_keeps_only_relevant(self):
        generated, retrieval, llm, _ = await self.run_graph(
            [("q1", [chunk(1, "eccentricity"), chunk(2, "cafeteria")]), ("directrix", [chunk(3, "directrix x=a/e")])],
            ['{"relevant":[1],"answerable":"partial","next_query":"directrix"}',
             '{"relevant":[1,2],"answerable":"yes","next_query":""}'],
        )
        second = retrieval.retrieve_for_bot.await_args_list[1].kwargs
        self.assertEqual(second["query"], "directrix")
        self.assertEqual(second["exclude_ids"], {"c1", "c2"})
        self.assertEqual([c.id for c in generated["retrieval"].chunks], ["c1", "c3"])
        self.assertEqual([r["decision"] for r in generated["rag_trace"]], ["refine", "done"])
        grader_call = llm.chat.await_args_list[0].kwargs
        self.assertEqual(grader_call["reasoning_effort"], "none")

    async def test_round_limit_stops_refinement(self):
        generated, retrieval, _, _ = await self.run_graph(
            [("q1", [chunk(1, "partial info")])],
            ['{"relevant":[1],"answerable":"partial","next_query":"more"}'],
            max_rounds=1,
        )
        self.assertEqual(retrieval.retrieve_for_bot.await_count, 1)
        self.assertEqual([c.id for c in generated["retrieval"].chunks], ["c1"])

    async def test_unparseable_grade_fails_open(self):
        generated, _, _, _ = await self.run_graph(
            [("q1", [chunk(1, "a"), chunk(2, "b")])], ["no idea"],
        )
        self.assertEqual([c.id for c in generated["retrieval"].chunks], ["c1", "c2"])

    async def test_irrelevant_evidence_falls_back_to_web(self):
        generated, _, _, tools = await self.run_graph(
            [("q1", [chunk(1, "cafeteria", score=0.5)])],
            ['{"relevant":[],"answerable":"no","next_query":""}'], web=True,
        )
        tools.execute_tool.assert_awaited_once()
        self.assertEqual(generated["route"], "tool")


if __name__ == "__main__":
    unittest.main()


class ProcessDocumentTests(unittest.IsolatedAsyncioTestCase):
    async def process(self, sections_batches, *, fail_on_upsert=False):
        from app.services import document_service as module

        document = SimpleNamespace(
            id=uuid.uuid4(), knowledge_base_id=KB, user_id=USER, storage_key="k",
            original_name="policy.pdf", extraction_metadata={}, status="uploaded", chunk_count=0,
        )
        kb = SimpleNamespace(chunking_config={"chunk_size": 120, "chunk_overlap": 20})
        store = SimpleNamespace(
            upsert_chunks=AsyncMock(side_effect=RuntimeError("qdrant down") if fail_on_upsert else (lambda **kw: len(kw["chunks"]))),
            delete_stale_ingests=AsyncMock(), delete_ingest=AsyncMock(),
        )

        async def iter_batches(path, *, stats):
            stats.source_type, stats.page_count = "pdf", len(sections_batches)
            for batch in sections_batches:
                stats.pages_with_text += 1
                yield batch

        async def update_metadata(db, doc, metadata):
            doc.extraction_metadata = metadata

        async def mark(status):
            async def _mark(db, doc, **kw):
                doc.status = status
                doc.chunk_count = kw.get("chunk_count", doc.chunk_count)
            return _mark

        service = module.DocumentService.__new__(module.DocumentService)
        service.extractor = SimpleNamespace(iter_batches=iter_batches)
        service.vector_store = store
        service.storage = SimpleNamespace(get_local_path=AsyncMock(return_value="/tmp/x.pdf"))
        db = SimpleNamespace(commit=AsyncMock(), rollback=AsyncMock(), refresh=AsyncMock())
        repo = module.DocumentRepository
        with patch.object(repo, "get_by_id", AsyncMock(return_value=document)), \
             patch.object(module.KnowledgeRepository, "get_by_id", AsyncMock(return_value=kb)), \
             patch.object(repo, "mark_processing", await mark("processing")), \
             patch.object(repo, "mark_ready", await mark("ready")), \
             patch.object(repo, "mark_failed", await mark("failed")), \
             patch.object(repo, "update_extraction_metadata", update_metadata):
            try:
                await service.process_document(db, document_id=document.id)
            except Exception as exc:
                document.error = exc
        return document, store

    def batch(self, page):
        return [ExtractedSection(text=f"Rule {page}. " + "Detail sentence here. " * 8, page=page,
                                 metadata={"section": "Policy"})]

    async def test_streams_batches_and_swaps_ingest(self):
        document, store = await self.process([self.batch(1), self.batch(2), self.batch(3)])
        self.assertEqual(document.status, "ready")
        calls = store.upsert_chunks.await_args_list
        self.assertGreater(len(calls), 1)
        ingest_ids = {c.kwargs["ingest_id"] for c in calls}
        self.assertEqual(len(ingest_ids), 1)
        indexes = [ch["index"] for c in calls for ch in c.kwargs["chunks"]]
        self.assertEqual(indexes, list(range(len(indexes))))
        self.assertEqual(document.chunk_count, len(indexes))
        self.assertEqual(calls[0].kwargs["chunks"][0]["metadata"]["file_name"], "policy.pdf")
        store.delete_stale_ingests.assert_awaited_once_with(document.id, keep_ingest_id=ingest_ids.pop())
        self.assertEqual(document.extraction_metadata["page_count"], 3)
        store.delete_ingest.assert_not_awaited()

    async def test_failure_removes_partial_ingest_and_marks_failed(self):
        document, store = await self.process([self.batch(1)], fail_on_upsert=True)
        self.assertEqual(document.status, "failed")
        self.assertIsInstance(document.error, RuntimeError)
        store.delete_ingest.assert_awaited_once()
        store.delete_stale_ingests.assert_not_awaited()

    async def test_empty_document_fails(self):
        document, store = await self.process([])
        self.assertEqual(document.status, "failed")
        store.delete_stale_ingests.assert_not_awaited()
